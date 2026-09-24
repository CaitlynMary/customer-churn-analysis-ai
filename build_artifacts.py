import nbformat as nbf
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json
import os

def create_notebook():
    nb = nbf.v4.new_notebook()
    
    cells = []
    
    cells.append(nbf.v4.new_markdown_cell("# Customer Churn Analysis and Prediction Using Data Analytics and Machine Learning\n\n## 1. Project Title\n**Customer Churn Analysis and Prediction Using Data Analytics and Machine Learning**\n\n## 2. Project Overview\nThis project analyzes customer data to identify patterns associated with churn and builds a machine learning model to predict whether a customer is likely to leave."))
    cells.append(nbf.v4.new_markdown_cell("## 3. Problem Statement\nCustomer churn is a critical metric for businesses. Identifying factors that lead to customer attrition allows companies to proactively address issues and improve retention rates."))
    cells.append(nbf.v4.new_markdown_cell("## 4. Objectives\n- Perform Exploratory Data Analysis (EDA) on customer data.\n- Identify key features influencing churn.\n- Build and evaluate a Machine Learning model (Logistic Regression) to predict churn."))
    cells.append(nbf.v4.new_markdown_cell("## 5. Dataset Description\nThe dataset used is the Telco Customer Churn dataset, containing demographics, services, account information, and churn status."))
    cells.append(nbf.v4.new_markdown_cell("## 6. Import Required Libraries"))
    cells.append(nbf.v4.new_code_cell("import pandas as pd\nimport numpy as np\nimport matplotlib.pyplot as plt\nimport seaborn as sns\nfrom sklearn.model_selection import train_test_split\nfrom sklearn.preprocessing import StandardScaler, OneHotEncoder\nfrom sklearn.compose import ColumnTransformer\nfrom sklearn.pipeline import Pipeline\nfrom sklearn.linear_model import LogisticRegression\nfrom sklearn.metrics import accuracy_score, precision_score, recall_score, f1_score, roc_auc_score, roc_curve, confusion_matrix, classification_report\n\nimport warnings\nwarnings.filterwarnings('ignore')"))
    
    cells.append(nbf.v4.new_markdown_cell("## 7. Load Dataset"))
    cells.append(nbf.v4.new_code_cell("df = pd.read_csv('../data/WA_Fn-UseC_-Telco-Customer-Churn.csv')\ndf.head()"))
    
    cells.append(nbf.v4.new_markdown_cell("## 8. Data Understanding"))
    cells.append(nbf.v4.new_code_cell("print(f\"Shape of dataset: {df.shape}\")\ndf.info()"))
    cells.append(nbf.v4.new_code_cell("df.describe()"))
    
    cells.append(nbf.v4.new_markdown_cell("## 9. Data Cleaning"))
    cells.append(nbf.v4.new_code_cell("# Handle empty spaces in TotalCharges\ndf['TotalCharges'] = df['TotalCharges'].replace(\" \", np.nan)\ndf['TotalCharges'] = pd.to_numeric(df['TotalCharges'])\n\n# Drop rows with missing TotalCharges\ndf = df.dropna(subset=['TotalCharges'])\nprint(f\"Dataset shape after cleaning: {df.shape}\")\n\n# Drop customerID\ndf = df.drop('customerID', axis=1)"))
    
    cells.append(nbf.v4.new_markdown_cell("## 10. Exploratory Data Analysis & 11. Data Visualization"))
    cells.append(nbf.v4.new_code_cell("sns.set_theme(style=\"whitegrid\")\n\n# Churn distribution\nplt.figure(figsize=(6,4))\nsns.countplot(data=df, x='Churn', palette='Set2')\nplt.title('Churn Distribution')\nplt.show()"))
    
    cells.append(nbf.v4.new_code_cell("# Churn by Contract Type\nplt.figure(figsize=(8,5))\nsns.countplot(data=df, x='Contract', hue='Churn', palette='Set2')\nplt.title('Churn by Contract Type')\nplt.show()"))
    
    cells.append(nbf.v4.new_code_cell("# Monthly Charges Distribution by Churn\nplt.figure(figsize=(8,5))\nsns.histplot(data=df, x='MonthlyCharges', hue='Churn', multiple=\"stack\", palette='Set2', bins=30)\nplt.title('Monthly Charges Distribution by Churn')\nplt.show()"))
    
    cells.append(nbf.v4.new_markdown_cell("## 12. Feature Engineering & 13. Data Preprocessing"))
    cells.append(nbf.v4.new_code_cell("X = df.drop('Churn', axis=1)\ny = df['Churn'].map({'Yes': 1, 'No': 0})\n\ncat_cols = X.select_dtypes(include=['object']).columns.tolist()\nnum_cols = X.select_dtypes(include=['number']).columns.tolist()\n\npreprocessor = ColumnTransformer(\n    transformers=[\n        ('num', StandardScaler(), num_cols),\n        ('cat', OneHotEncoder(drop='first', sparse_output=False), cat_cols)\n    ]\n)"))
    
    cells.append(nbf.v4.new_markdown_cell("## 14. Train-Test Split"))
    cells.append(nbf.v4.new_code_cell("X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42, stratify=y)\nprint(f\"Training data shape: {X_train.shape}\")\nprint(f\"Testing data shape: {X_test.shape}\")"))
    
    cells.append(nbf.v4.new_markdown_cell("## 15. Machine Learning Model"))
    cells.append(nbf.v4.new_code_cell("lr_model = Pipeline(steps=[\n    ('preprocessor', preprocessor),\n    ('classifier', LogisticRegression(random_state=42, max_iter=1000))\n])\n\nlr_model.fit(X_train, y_train)"))
    
    cells.append(nbf.v4.new_markdown_cell("## 16. Model Evaluation"))
    cells.append(nbf.v4.new_code_cell("y_pred = lr_model.predict(X_test)\ny_pred_prob = lr_model.predict_proba(X_test)[:, 1]\n\nprint(\"Accuracy:\", accuracy_score(y_test, y_pred))\nprint(\"Precision:\", precision_score(y_test, y_pred))\nprint(\"Recall:\", recall_score(y_test, y_pred))\nprint(\"F1-score:\", f1_score(y_test, y_pred))\nprint(\"ROC-AUC:\", roc_auc_score(y_test, y_pred_prob))\n\nprint(\"\\nClassification Report:\")\nprint(classification_report(y_test, y_pred))"))
    cells.append(nbf.v4.new_code_cell("# Confusion Matrix\ncm = confusion_matrix(y_test, y_pred)\nplt.figure(figsize=(6,4))\nsns.heatmap(cm, annot=True, fmt='d', cmap='Blues', xticklabels=['No Churn', 'Churn'], yticklabels=['No Churn', 'Churn'])\nplt.title('Confusion Matrix - Logistic Regression')\nplt.show()"))
    
    cells.append(nbf.v4.new_code_cell("# ROC Curve\nfpr, tpr, thresholds = roc_curve(y_test, y_pred_prob)\nplt.figure(figsize=(6,4))\nplt.plot(fpr, tpr, label=f'Logistic Regression (AUC = {roc_auc_score(y_test, y_pred_prob):.4f})')\nplt.plot([0, 1], [0, 1], 'k--')\nplt.title('ROC Curve')\nplt.xlabel('False Positive Rate')\nplt.ylabel('True Positive Rate')\nplt.legend(loc='lower right')\nplt.show()"))
    
    cells.append(nbf.v4.new_markdown_cell("## 17. Feature Importance / Interpretation"))
    cells.append(nbf.v4.new_code_cell("feature_names = num_cols + list(lr_model.named_steps['preprocessor'].named_transformers_['cat'].get_feature_names_out(cat_cols))\ncoefficients = lr_model.named_steps['classifier'].coef_[0]\n\nfeat_imp = pd.DataFrame({'Feature': feature_names, 'Coefficient': coefficients})\nfeat_imp['Abs_Coefficient'] = feat_imp['Coefficient'].abs()\nfeat_imp = feat_imp.sort_values(by='Abs_Coefficient', ascending=False).head(10)\n\nplt.figure(figsize=(10,6))\nsns.barplot(data=feat_imp, x='Coefficient', y='Feature', palette='coolwarm')\nplt.title('Top 10 Feature Coefficients (Logistic Regression)')\nplt.show()"))
    
    cells.append(nbf.v4.new_markdown_cell("## 18. Example Churn Prediction"))
    cells.append(nbf.v4.new_code_cell("example_data = pd.DataFrame([X_test.iloc[0]])\nexample_pred = lr_model.predict(example_data)\nexample_prob = lr_model.predict_proba(example_data)[:, 1]\n\nprint(f\"Input features:\\n{example_data.to_dict(orient='records')[0]}\")\nprint(f\"Predicted Churn: {'Yes' if example_pred[0] == 1 else 'No'}\")\nprint(f\"Churn Probability: {example_prob[0]:.4f}\")"))
    
    cells.append(nbf.v4.new_markdown_cell("## 19. Key Insights\n- Customers with Month-to-month contracts are significantly more likely to churn compared to one-year or two-year contracts.\n- Higher monthly charges are generally associated with a higher likelihood of churn.\n- Certain payment methods, specifically electronic check, are associated with a higher churn rate.\n- Tenure is negatively correlated with churn (longer tenure = less likely to churn)."))
    
    cells.append(nbf.v4.new_markdown_cell("## 20. Conclusion\nThe Logistic Regression model successfully predicts customer churn with reasonable accuracy. The analysis reveals that contract type, tenure, and monthly charges are strong predictors of churn. Businesses can use these insights to target high-risk customers with retention offers, such as incentives to switch to longer-term contracts."))
    
    cells.append(nbf.v4.new_markdown_cell("## 21. Future Enhancements\n- Hyperparameter tuning using GridSearchCV.\n- Experimenting with tree-based models like Random Forest and XGBoost.\n- Deploying the model as a REST API using Flask or FastAPI.\n- Building an interactive dashboard for real-time churn monitoring."))
    
    nb['cells'] = cells
    with open('notebooks/Customer_Churn_Analysis_and_Prediction.ipynb', 'w') as f:
        nbf.write(nb, f)
    print("Jupyter Notebook created at notebooks/Customer_Churn_Analysis_and_Prediction.ipynb")

def create_report():
    if not os.path.exists('outputs/metrics.json'):
        print("metrics.json not found! Please run the analysis script first.")
        return
        
    with open('outputs/metrics.json', 'r') as f:
        metrics = json.load(f)
        
    doc = Document()
    
    # 1. Cover Page
    doc.add_heading('Project Report', 0)
    doc.add_paragraph('Customer Churn Analysis and Prediction Using Data Analytics and Machine Learning')
    doc.add_paragraph('AICTE | IBM SkillsBuild Data Analytics with AI Academic Internship 2026')
    doc.add_paragraph('Conducted by BharatCares in association with AICTE')
    doc.add_paragraph('\nSubmitted by: J S Caitlyn Mary')
    doc.add_page_break()
    
    # Rest of the report
    sections = [
        "2. Declaration",
        "This is to certify that this project is submitted as part of the AICTE | IBM SkillsBuild Data Analytics with AI Academic Internship 2026.",
        
        "3. Acknowledgement",
        "I would like to thank AICTE, IBM, and BharatCares for providing this internship opportunity.",
        
        "4. Abstract",
        "Customer churn is a critical issue for telecommunications companies. This project explores the Telco Customer Churn dataset, visualizes key patterns, and builds a Logistic Regression model to predict customer churn based on historical data.",
        
        "5. Table of Contents",
        "1. Project Report (Cover)\n2. Declaration\n3. Acknowledgement\n4. Abstract\n5. Table of Contents\n6. Introduction\n7. Problem Statement\n8. Objectives\n9. Scope of the Project\n10. Dataset Description\n11. Technologies Used\n12. System/Project Methodology\n13. Data Preprocessing\n14. Exploratory Data Analysis & 15. Data Visualization\n16. Machine Learning Methodology\n17. Model Training\n18. Model Evaluation\n19. Results\n20. Key Findings\n21. Conclusion\n22. Future Enhancements\n23. References",
        
        "6. Introduction",
        "Customer churn refers to when a customer ceases their relationship with a company. Predicting churn allows companies to intervene proactively.",
        
        "7. Problem Statement",
        "Identify the demographic and service-related factors associated with customer churn and develop a machine learning model to predict it accurately.",
        
        "8. Objectives",
        "- Analyze customer attributes to find churn associations.\n- Preprocess data for machine learning.\n- Train and evaluate a Logistic Regression model.",
        
        "9. Scope of the Project",
        "The project covers Data Cleaning, Exploratory Data Analysis (EDA), Feature Engineering, Model Training, and Evaluation using Python and scikit-learn.",
        
        "10. Dataset Description",
        "The dataset used is the Telco Customer Churn dataset. It includes variables such as Tenure, MonthlyCharges, TotalCharges, Contract type, and PaymentMethod.",
        
        "11. Technologies Used",
        "- Python\n- Pandas, NumPy\n- Matplotlib, Seaborn\n- Scikit-learn\n- Jupyter Notebook",
        
        "12. System/Project Methodology",
        "The project follows the standard data science lifecycle: Data Collection -> Data Cleaning -> EDA -> Preprocessing -> Model Training -> Evaluation.",
        
        "13. Data Preprocessing",
        "Categorical variables were encoded using OneHotEncoder. Numerical variables were scaled using StandardScaler. Missing values in TotalCharges were removed.",
        
        "14. Exploratory Data Analysis & 15. Data Visualization",
        "During EDA, it was observed that month-to-month contracts and higher monthly charges are associated with higher churn.",
        
        "16. Machine Learning Methodology",
        "Logistic Regression was selected due to its interpretability. A Pipeline was constructed to chain preprocessing and model training.",
        
        "17. Model Training",
        "The dataset was split into 80% training and 20% testing sets using stratified sampling to maintain the class distribution.",
        
        "18. Model Evaluation",
        "The model was evaluated on the test set. Key metrics are reported in the Results section.",
        
        "19. Results",
        f"- Accuracy: {metrics['Accuracy']:.4f}\n" +
        f"- Precision: {metrics['Precision']:.4f}\n" +
        f"- Recall: {metrics['Recall']:.4f}\n" +
        f"- F1-Score: {metrics['F1_Score']:.4f}\n" +
        f"- ROC-AUC: {metrics['ROC_AUC']:.4f}",
        
        "20. Key Findings",
        "- Month-to-month contracts are highly associated with customer churn.\n- Customers with longer tenure are less likely to churn.\n- High monthly charges correspond to a higher risk of churn.",
        
        "21. Conclusion",
        f"The Logistic Regression model achieved an accuracy of {metrics['Accuracy']*100:.2f}%, precision of {metrics['Precision']*100:.2f}%, recall of {metrics['Recall']*100:.2f}%, F1-score of {metrics['F1_Score']*100:.2f}%, and ROC-AUC of {metrics['ROC_AUC']*100:.2f}% on the test dataset.",
        
        "22. Future Enhancements",
        "- Hyperparameter tuning for improved recall.\n- Implementation of ensemble models like Random Forest.\n- Creation of an interactive web dashboard using Flask.",
        
        "23. References",
        "- Scikit-learn Documentation: https://scikit-learn.org/\n- Telco Customer Churn Dataset Source: https://raw.githubusercontent.com/IBM/telco-customer-churn-on-icp4d/master/data/Telco-Customer-Churn.csv"
    ]
    
    for idx in range(0, len(sections), 2):
        doc.add_heading(sections[idx], level=1)
        doc.add_paragraph(sections[idx+1])
        
        # Embed corresponding charts after specific sections
        if "14. Exploratory Data Analysis & 15. Data Visualization" in sections[idx]:
            if os.path.exists('outputs/churn_distribution.png'):
                doc.add_picture('outputs/churn_distribution.png', width=Inches(5))
                doc.add_paragraph('Figure 1: Customer Churn Distribution')
            if os.path.exists('outputs/churn_by_contract.png'):
                doc.add_picture('outputs/churn_by_contract.png', width=Inches(5))
                doc.add_paragraph('Figure 2: Customer Churn by Contract Type')
            if os.path.exists('outputs/monthly_charges_dist.png'):
                doc.add_picture('outputs/monthly_charges_dist.png', width=Inches(5))
                doc.add_paragraph('Figure 3: Monthly Charges Distribution by Churn')
        elif "18. Model Evaluation" in sections[idx]:
            if os.path.exists('outputs/confusion_matrix.png'):
                doc.add_picture('outputs/confusion_matrix.png', width=Inches(4))
                doc.add_paragraph('Figure 4: Confusion Matrix (Logistic Regression)')
            if os.path.exists('outputs/roc_curve.png'):
                doc.add_picture('outputs/roc_curve.png', width=Inches(4))
                doc.add_paragraph('Figure 5: ROC Curve (Logistic Regression)')
        elif "20. Key Findings" in sections[idx]:
            if os.path.exists('outputs/feature_importance.png'):
                doc.add_picture('outputs/feature_importance.png', width=Inches(5.5))
                doc.add_paragraph('Figure 6: Top 10 Feature Importances (Coefficients)')
        
    doc.save('Customer_Churn_ProjectReport.docx')
    print("Report created at Customer_Churn_ProjectReport.docx")

if __name__ == "__main__":
    # create_notebook() # Disabled to avoid modifying the notebook as requested
    create_report()
